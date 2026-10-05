from pathlib import Path
from datetime import datetime
import shutil
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

WORKSPACE = Path(__file__).resolve().parent.parent
ROOT = WORKSPACE / 'docs'
BACKUP = WORKSPACE / 'archive' / 'originals_before_competition_rewrite_20261004'
BACKUP.mkdir(exist_ok=True)
for lang in ('ZH', 'RU', 'EN'):
    source = ROOT / f'SmartMed_Ambient_EHR_Copilot_{lang}.docx'
    if not (BACKUP / source.name).exists():
        shutil.copy2(source, BACKUP / source.name)

# All three versions share the same page and section structure. Statements below
# describe the proposed implementation, not unperformed experiments.
CONTENT = {'ZH': {'title': '智录门诊病历助手',
        'subtitle': 'SmartMed 技术方案',
        'meta': '第八届全球校园人工智能算法精英大赛\n算法创新赛道 赛题4 AI+场景创新\n2026年10月5日',
        'pages': [[('h', '1 项目概述'),
                   ('h2', '作品简介'),
                   ('p',
                    '本项目面向中文门诊沟通后的病历整理场景，拟将医患对话转为可追溯的结构化病历草稿。系统结合中文语音识别、现成大语言模型、术语检索与确定性规则，为关键字段关联原文证据，对缺失信息保留空值，并提示过敏、重复成分及用药信息不完整等问题。医生核对后方可导出。首阶段拟完成本地运行的中文演示与独立测试，验证记录完整性、事实一致性和流程可靠性；临床效果尚待后续验证。'),
                   ('h2', '场景与用户需求'),
                   ('p', '目标用户是在门诊完成沟通后整理病历的医生。对话中的症状、时长、否认项、既往史和药物信息分散在多个轮次，整理时需要反复查找原文。系统拟提供同屏对照：左侧显示对话与音频片段，右侧显示结构化草稿、证据定位和待确认事项。'),
                   ('p', '首阶段限定为模拟中文门诊记录，先验证文本输入，再接入短音频文件。主观资料、客观资料、医生已表达的评估以及计划按SOAP结构组织。未在输入中出现的诊断、检查结果和治疗方案不得由模型补写。'),
                   ('h2', '阶段目标与实施边界'),
                   ('p',
                    '截至10月5日，已实现本地Qwen3.5 '
                    '4B文本抽取、Instructor与Pydantic校验、SQLite版本记录、编辑确认及JSON导出。已比较四个Qwen模型，并对已知错误进行回归测试。当前接口提取六项事实；完整SOAP、语音界面、术语检索及用药检查仍待实现。尚无临床试点或医学审核结果。'),
                   ('h2', '团队能力与验证边界'),
                   ('p', '团队由两名软件开发成员组成。成员A负责后端、模型、字段契约、验证与存储；成员B负责音频、界面、数据整理和演示材料，双方交叉核对案例。团队目前没有医生或临床顾问。首阶段由开发成员操作模拟场景，不能将确认按钮或开发者复核宣称为医生验证。')],
                  [('h', '2 系统架构与处理流程'),
                   ('p', '采用单一中文网页与一个后端服务。前端拟使用React和TypeScript，后端拟使用FastAPI，SQLite保存任务、草稿版本和确认记录；音频及数据文件保存在本地。后端按明确状态组织流程，模型通过适配器接入，避免在首阶段增加分布式部署负担。'),
                   ('h2', '从输入到可核对草稿'),
                   ('p', '中文对话或音频文件 → 转写及人工校正 → 结构化抽取 → 字段与原文关联 → 术语候选检索 → 规则检查 → 医生编辑确认 → 导出。文本输入可直接进入抽取阶段；音频需先展示转写供用户校正。任何环节失败均保留当前任务并显示具体原因。'),
                   ('p', '语音模块拟先测试FunASR中文模型。说话人识别为可选项；如果无法可靠区分医生与患者，界面要求人工标注，不能把自动分离作为必达指标。演示所需模型与词表预先下载，以验证断网运行路径。'),
                   ('p',
                    '当前抽取模型为本地量化Qwen3.5 4B，关闭思考输出。以Qwen2.5 3B、Qwen3 4B Instruct和Qwen2.5 7B作对照。Instructor传递JSON '
                    'Schema；按状态约束值与证据，证据限于患者原文片段，并再次验证字面一致性。最多重试一次；失败返回错误。结构合规不等于医学正确。'),
                   ('h2', '数据契约与医生控制'),
                   ('p',
                    '当前六项为发热、咳嗽、症状持续时间、青霉素过敏、已用药物及已用药物剂量。状态为present、negated、unknown；未解决的冲突归入unknown。值采用原文片段，不擅自规范化数字与单位；否认也须引用原话。医生是目标用户角色，当前演示由开发成员在模拟场景中操作。'),
                   ('p', '诊断术语由医生已表达的内容触发检索，并展示词表来源和版本。没有可核实的编码来源时仅展示术语候选，不生成虚假ICD编码。首阶段规则覆盖已知过敏与成分匹配、重复成分及剂量信息缺失；未覆盖或信息不足显示“未知”，不能显示“安全”。'),
                   ('p', '医生可编辑草稿、查看证据和处理提示。确认后导出JSON与一种可打印病历格式，记录草稿版本和确认时间。该确认属于原型操作记录，不宣称具备法律电子签名效力。')],
                  [('h', '3 数据方案与来源'),
                   ('h2', '先选择来源 再补充案例'),
                   ('p', '数据分为对话样本、结构化标注、语音测试及术语规则四类。先检查来源、许可、内容质量和可下载性，再确定使用范围；不直接将大规模医疗问答当作SOAP训练集。每条样本记录来源、原始编号、版本、处理方式与分组，禁止同源改写样本跨开发集和测试集。'),
                   ('p', '优先检查IMCS-21。该研究提供中文医疗咨询与报告生成任务，适合观察对话到报告的映射，但主要覆盖儿科疾病，报告结构也不等于本项目SOAP契约。选用前需核对数据使用条件，并人工建立字段与证据标注。外部结果不得作为本项目成绩。'),
                   ('p', 'DISC-Med-SFT作为补充候选，其数据包含加工重构的咨询内容；Huatuo-26M主要为医疗问答，适合术语和场景参考。两者不能直接作为真实录音、独立人工金标准或本项目临床证据。代码许可不能替代所有原始数据的使用授权。'),
                   ('h2', '首阶段小规模评测集'),
                   ('p', '四模型比较已使用40个AI辅助编写的中文模拟对话，共240个字段；未使用真实患者数据。案例及参考答案未经医生或中文母语者审核。此集用于模型选择，随后部分案例用于提示与验证器开发，不能再视为独立终测。下一步另建未参与开发的冻结测试集，保留来源及审核记录。'),
                   ('p', '对照标注至少覆盖症状、时间、否认项、过敏、药物名称与明确出现的剂量，附来源轮次。两名成员交叉复核；分歧单独记录。具备条件时邀请中文熟练者核对语言，临床解释仍需医学背景人员验证，不将机器翻译视为医学审查。'),
                   ('h2', '音频与规则数据'),
                   ('p', '拟准备6段1至2分钟中文模拟音频，开发与测试各3段，保留人工转写和角色标注。合成语音与人工朗读分别记录，不宣称代表真实门诊声学条件。通用语音库可作辅助测试，但不能替代医疗用语测试。'),
                   ('p', '药物成分、过敏映射和诊断术语使用小规模、可溯源、带版本的词表。每条规则记录来源、适用条件及覆盖范围；没有可靠依据时不增加剂量上限或复杂相互作用判断。首阶段不采集真实患者录音。')],
                  [('h', '4 创新点与验证方法'),
                   ('h2', '可验证的场景贡献'),
                   ('p', '第一，病历字段与对话证据同屏关联，使使用者能逐项追溯。第二，将否认、未知与矛盾作为显式状态，减少模型在信息缺失时自动补写。第三，将术语检索和确定性检查放在模型输出之后，并把覆盖范围与医生确认纳入统一流程。'),
                   ('p', '上述内容是拟验证的工程与场景贡献。方案不主张原创基础模型，也不声称已优于成熟临床产品。模型、工具和数据来源单独列明；团队贡献集中在流程、数据契约、界面、规则及评测。'),
                   ('h2', '对比与消融设计'),
                   ('p',
                    '首次比较使用相同40例、六字段、提示和配置。Qwen3.5 4B状态正确235/240，结构校验通过28/40；Qwen3 4B '
                    'Instruct为230/240、37/40。Qwen3.5的14次暖机请求中位数为7.73秒；该配置不等同当前服务。改进后五个已知案例的状态和值均匹配参考，53项后端自动测试通过；回归结果不能证明临床可靠性或独立测试提升。后续在新冻结集进行对比与消融。'),
                   ('p', '语音测试分别评估人工转写输入与ASR转写输入，以区分抽取错误和语音错误。记录逐案例错误、时延与资源使用，不仅报告平均值。模拟集规模有限，评测结论仅适用于所测场景。'),
                   ('table',
                    [['验证项', '计算或检查方法'],
                     ['字段抽取', '按预先定义的字段和值匹配计算P、R、F1，并单独报告否认项和剂量。'],
                     ['事实与证据', '人工核对无依据事实数量、证据准确率以及缺失字段是否保留未知。'],
                     ['词表与规则', '在已覆盖条目上统计候选命中与提示正确性；另报未覆盖和信息不足。'],
                     ['语音与流程', '人工转写计算CER；仅有角色金标准时计算DER；检查确认和导出路径。']]),
                   ('h2', '验收门槛'),
                   ('p', '正常记录、过敏提示和信息缺失三个中文场景须完成端到端演示；无法解析的输出不能进入确认状态；输入未提及的剂量在专项测试中须保留未知；未确认草稿不能作为已确认病历导出。质量指标在实验后据实填写，目前不承诺特定F1、速度或节时比例。')],
                  [('h', '5 首阶段实施与演示'),
                   ('p', '开发窗口为2026年10月4日至9日，围绕一个可运行闭环交付。校赛完成节点按官方通知为10月10日，具体校内提交时间仍以学校通知为准；区域赛报名及缴费截止为10月15日20时。后续区域赛与国赛准备另行迭代，不把12月作为当前比赛交付节点。'),
                   ('table',
                    [['时间', '成员A主要工作', '成员B主要工作', '交付'],
                     ['10月4日至5日', '字段契约、文本API及存储', '数据来源、案例及界面骨架', '可运行骨架与版本化数据'],
                     ['10月6日', '模型、校验、证据关联及API', '音频与ASR；界面', '文本草稿与音频输入'],
                     ['10月7日', '存储、确认、导出及联调', '证据与音频核对界面；联调', '三类场景闭环'],
                     ['10月8日', '运行对比实验并编写实际结果', '复测流程、失败恢复与断网演示', '结果表与错误清单'],
                     ['10月9日', '整理中文方案、答辩内容与提交包', '录制中文演示并核验打包运行', '演示视频与提交材料']]),
                   ('h2', '演示安排'),
                   ('p', '视频目标时长约4分钟。先用正常案例展示中文输入、病历字段和原文证据，再用过敏案例展示规则提示，最后用信息缺失案例展示未知状态和医生补充。结尾展示确认、导出与真实评测结果；不以剪辑掩盖模型失败或时延。'),
                   ('h2', '交付内容与资源'),
                   ('p', '提交材料包括中文技术方案PDF、中文演示MP4、答辩PPT导出的PDF，以及代码、说明、模型版本、数据清单和评测脚本。根据实际情况提供真实佐证材料，不编造医院合作、专利或试点证明。'),
                   ('p', '本地资源先按16GB内存设备进行顺序任务测试，记录模型载入后的峰值内存与时延。必要时缩小模型和输入长度；首阶段不默认购买GPU或租用云训练资源。运行说明须列出软件版本、模型下载步骤和失败处理方法。')],
                  [('h', '6 应用价值与发展计划'),
                   ('p', '预期价值是帮助医生整理沟通内容并核对记录依据。是否节约时间、减少遗漏及适合临床使用，需要通过后续用户测试和医学验证确认；当前模拟评测不能替代临床试验，也不能证明诊断或用药安全。'),
                   ('p', '推广路径拟从教学或模拟记录工具开始，经过中文医学人员审核和受控试用，再讨论医院系统接口。商业形态可探索部署及维护服务，但收费、采购和收益均未得到验证。'),
                   ('h2', '主要风险与迭代条件'),
                   ('p', '中文医学表达与团队语言能力存在差距，需保留原文、交叉复核并寻求中文审核。ASR错误可能改变否认词或剂量，因此关键字段必须可回看；规则覆盖有限，明确显示未知。数据许可不明的来源不进入正式数据包，未来真实数据需要另外建立授权和保护流程。'),
                   ('p', '首阶段后按错误频率决定改进顺序：先修正数据与提示，再评估检索扩展；只有在足够数量、许可明确的对话与字段标注可用、基线已完成且算力满足时，才考虑微调。更大规模测试、可靠说话人识别和真实业务接口列为后续工作。'),
                   ('h2', '参考来源'),
                   ('ref', '赛事通知与材料规范\nhttps://www.aicomp.cn/notice/notice-1/3674.html\nhttps://www.aicomp.cn/tracks/tracks-2/3775.html'),
                   ('ref', 'IMCS-21 数据与研究实现\nhttps://github.com/lemuria-wchen/imcs21'),
                   ('ref', '补充数据候选 DISC-MedLLM 与 Huatuo-26M\nhttps://github.com/FudanDISC/DISC-MedLLM\nhttps://github.com/FreedomIntelligence/Huatuo-26M'),
                   ('ref', '当前模型与拟测试语音工具\nhttps://huggingface.co/Qwen/Qwen3.5-4B\nhttps://github.com/modelscope/FunASR'),
                   ('p', '参考来源检查日期为2026年10月4日。数据是否纳入项目以实际许可核查和样本检查结果为准。')]]},
 'RU': {'title': 'Ассистент подготовки амбулаторной записи',
        'subtitle': 'SmartMed Техническое предложение',
        'meta': 'Восьмой Global Campus Artificial Intelligence Algorithm Elite Competition\nАлгоритмический трек Тема 4 AI+场景创新\n5 октября 2026 года',
        'pages': [[('h', '1 Обзор проекта'),
                   ('h2', 'Описание для заявки'),
                   ('p',
                    'SmartMed предназначен для подготовки проверяемого черновика медицинской записи по китайскому диалогу врача и пациента. Планируется '
                    'объединить распознавание речи, готовую языковую модель, поиск терминов и детерминированные проверки. Каждое существенное поле связывается '
                    'с исходной репликой; отсутствующие сведения остаются неизвестными. Система показывает возможную аллергию, повторяющиеся действующие '
                    'вещества и неполные сведения о лекарствах. Экспорт следует после проверки врачом. Первый этап включает локальную демонстрацию на '
                    'китайском и независимое тестирование. Клинический эффект пока не подтверждён.'),
                   ('h2', 'Сценарий и потребность'),
                   ('p',
                    'Пользователь — врач, оформляющий запись после приёма. Симптомы, длительность, отрицания, анамнез и сведения о лекарствах распределены по '
                    'репликам; при подготовке записи приходится искать их в разговоре. Предлагается общий экран: слева диалог и фрагменты аудио, справа поля '
                    'черновика, ссылки на доказательства и вопросы для уточнения.'),
                   ('p',
                    'Первый этап ограничен моделируемыми китайскими приёмами. Сначала проверяется текстовый ввод, затем короткие аудиофайлы. Структура SOAP '
                    'содержит субъективные сведения, объективные данные, уже высказанную врачом оценку и план. Модель не должна дописывать диагнозы, '
                    'результаты обследований или лечение, которых нет во входе.'),
                   ('h2', 'Цель первого этапа'),
                   ('p',
                    'На 5 октября реализованы локальное извлечение Qwen3.5 4B, проверки Instructor и Pydantic, версии SQLite, редактирование, подтверждение и '
                    'JSON-экспорт. Проведены сравнение четырёх Qwen и регрессия известных ошибок. API извлекает шесть фактов; полный SOAP, аудиоинтерфейс, '
                    'поиск терминов и проверки назначений ещё разрабатываются. Клинического пилота и медицинской проверки нет.'),
                   ('h2', 'Команда и границы проверки'),
                   ('p',
                    'В команде два разработчика. Участник A отвечает за бэкенд, модель, контракт, проверки и хранение. Участник B — за аудио, интерфейс, '
                    'примеры и материалы; случаи проверяются совместно. Врача или клинического консультанта в команде нет. Демо выполняют разработчики на '
                    'моделируемых случаях; подтверждение не считается проверкой врачом.')],
                  [('h', '2 Архитектура и обработка'),
                   ('p',
                    'Одна китайская веб-панель и один backend. Планируемый стек: React и TypeScript для интерфейса, FastAPI для сервиса, SQLite для задач, '
                    'версий черновика и подтверждений. Аудио и наборы данных хранятся локально. Backend ведёт явные состояния процесса, а модели подключаются '
                    'через адаптеры.'),
                   ('h2', 'Путь от ввода до черновика'),
                   ('p',
                    'Китайский текст или аудиофайл → транскрипция и исправление → извлечение полей → привязка к репликам → поиск терминов → проверки → '
                    'редактирование и подтверждение врачом → экспорт. Текст сразу поступает на извлечение. Для аудио сначала показывается распознанный текст. '
                    'Ошибка сохраняет текущую задачу и объясняется пользователю.'),
                   ('p',
                    'Для речи сначала проверяется китайская модель FunASR. Автоматические роли говорящих опциональны: при ненадёжном разделении интерфейс '
                    'требует ручной разметки. Для проверки работы без сети модели и словари загружаются заранее.'),
                   ('p',
                    'Текущая модель — локальная квантованная Qwen3.5 4B без вывода рассуждений. Контроли — Qwen2.5 3B, Qwen3 4B Instruct и Qwen2.5 7B. '
                    'Instructor передаёт JSON Schema с ограничениями статусов, значений и цитат. Цитата выбирается из исходных фрагментов пациента и сверяется '
                    'с оригиналом. Разрешён один повтор, затем возвращается ошибка. Валидная структура не доказывает медицинскую правильность.'),
                   ('h2', 'Поля и управление врачом'),
                   ('p',
                    'Текущие поля: температура, кашель, длительность, аллергия на пенициллин, принятые лекарства и доза. Статусы — present, negated, unknown; '
                    'неразрешённый конфликт относится к unknown. Значения копируются из оригинала без нормализации чисел и единиц; отрицание тоже требует '
                    'цитату. Врач обозначает целевую роль пользователя; в демо её выполняет разработчик.'),
                   ('p',
                    'Поиск диагнозов запускается для содержания, уже выраженного врачом. Кандидаты сопровождаются источником и версией словаря. Если нет '
                    'проверенного источника кодов, показываются термины без выдуманного ICD. Первые правила проверяют известную аллергию, повторение вещества '
                    'и неполноту дозировки. Вне покрытия и при недостатке данных результат — «неизвестно», а не «безопасно».'),
                   ('p',
                    'Врач редактирует поля, просматривает цитаты и разбирает предупреждения. После подтверждения доступны JSON и один печатный формат; '
                    'сохраняются версия и время. Это запись действия в прототипе, без заявления о юридической электронной подписи.')],
                  [('h', '3 Данные и происхождение'),
                   ('h2', 'Источники до генерации'),
                   ('p',
                    'Разделяются четыре вида данных: диалоги, эталонные поля, аудиотесты и справочники правил. До включения проверяются происхождение, '
                    'лицензия, доступность и качество. Медицинские вопросы и ответы не объявляются готовым обучающим набором SOAP. У каждого примера '
                    'фиксируются источник, исходный ID, версия, преобразование и группа; варианты одного исходника не попадают в разные выборки.'),
                   ('p',
                    'Первым кандидатом выбран IMCS-21: китайские консультации и задача генерации медицинского отчёта. Корпус преимущественно педиатрический; '
                    'его структура отличается от нашего SOAP. Требуются проверка условий использования и ручная привязка полей к репликам. Результаты авторов '
                    'корпуса нельзя представлять как результаты SmartMed.'),
                   ('p',
                    'DISC-Med-SFT — дополнительный кандидат с переработанными консультациями. Huatuo-26M преимущественно содержит медицинские вопросы и ответы '
                    'и может помочь с терминологией и сценариями. Эти данные не являются реальными записями приёма или независимым эталоном проекта. Лицензия '
                    'кода не гарантирует права на все исходные данные.'),
                   ('h2', 'Небольшой набор первого этапа'),
                   ('p',
                    'Сравнение четырёх моделей использует 40 синтетических китайских диалогов, составленных с помощью ИИ, и 240 полей. Реальных пациентов нет; '
                    'врач и носитель китайского эталоны не проверяли. Набор использован для выбора модели, затем часть случаев — для доработки подсказок и '
                    'проверок. Он не является независимым итоговым тестом. Далее нужен новый замороженный набор, не участвующий в разработке, с учётом '
                    'происхождения и проверки.'),
                   ('p',
                    'Разметка включает симптомы, время, отрицания, аллергию, лекарства и прямо названные дозы со ссылками на реплики. Оба участника проверяют '
                    'работу друг друга и фиксируют расхождения. При доступности привлекается человек со свободным китайским; клиническое толкование требует '
                    'медицинской проверки. Машинный перевод её не заменяет.'),
                   ('h2', 'Аудио и справочники'),
                   ('p',
                    'Планируются 6 китайских записей моделируемого приёма по 1–2 минуты: по 3 для разработки и теста, с ручной транскрипцией и ролями. Синтез '
                    'и чтение человеком обозначаются отдельно; такие записи не представляют реальные акустические условия клиники. Общие речевые корпуса '
                    'используются лишь дополнительно.'),
                   ('p',
                    'Состав лекарств, соответствия аллергий и термины диагнозов берутся из небольших проверяемых словарей с версиями. Для каждого правила '
                    'указываются источник, условия и покрытие. Без надёжного основания не добавляются предельные дозы и сложные взаимодействия. Реальные '
                    'записи пациентов на первом этапе не собираются.')],
                  [('h', '4 Вклад и проверка качества'),
                   ('h2', 'Проверяемые улучшения сценария'),
                   ('p',
                    'Первый вклад — связь поля записи с доказательством в диалоге на общем экране. Второй — явная обработка отрицаний, неизвестности и '
                    'конфликтов, предотвращающая заполнение пропусков по догадке. Третий — поиск терминов и детерминированные проверки после модели с показом '
                    'покрытия и подтверждением врачом.'),
                   ('p',
                    'Это планируемый инженерный вклад, который предстоит проверить. Оригинальность базовой модели и превосходство над клиническими продуктами '
                    'не заявляются. Чужие модели, инструменты и наборы цитируются; вклад команды относится к процессу, контрактам, интерфейсу, правилам и '
                    'оценке.'),
                   ('h2', 'Сравнение и анализ компонентов'),
                   ('p',
                    'В исходном сравнении одинаковы 40 случаев, шесть полей, подсказка и настройки. У Qwen3.5 4B верны 235/240 статусов, структуру проходят '
                    '28/40 ответов; у Qwen3 4B Instruct — 230/240 и 37/40. Медиана 14 тёплых запросов Qwen3.5 — 7,73 с; это другая конфигурация, чем текущий '
                    'API. После правок пять известных случаев совпали по статусам и значениям, прошли 53 теста бэкенда. Регрессия не доказывает клиническую '
                    'надёжность или прирост на независимых данных. Следующий эксперимент требует новой замороженной выборки.'),
                   ('p',
                    'На аудио сравниваются извлечение из ручной транскрипции и из результата ASR, чтобы разделить ошибки речи и извлечения. Сохраняются ошибки '
                    'каждого случая, задержки и расход ресурсов. Выводы небольшого моделируемого набора ограничены проверенными сценариями.'),
                   ('table',
                    [['Предмет оценки', 'Способ проверки'],
                     ['Извлечение', 'P, R и F1 по заранее определённым полям и значениям; отдельный результат для отрицаний и доз.'],
                     ['Факты и цитаты', 'Ручной подсчёт неподтверждённых фактов, верных цитат и правильно оставленных неизвестными полей.'],
                     ['Словарь и правила', 'Попадания кандидатов и корректность предупреждений в пределах покрытия; отдельно неизвестность.'],
                     ['Речь и процесс', 'CER по ручной транскрипции; DER только при эталонных ролях; проверка подтверждения и экспорта.']]),
                   ('h2', 'Приёмка'),
                   ('p',
                    'Три китайских сценария — обычная запись, аллергия и пропуск — проходят весь путь. Невалидный JSON не допускается к подтверждению; '
                    'отсутствующая доза остаётся неизвестной в специальных тестах; неподтверждённый черновик не экспортируется как подтверждённая запись. '
                    'Численные результаты заполняются после экспериментов. Заранее не обещаются F1, скорость или процент экономии времени.')],
                  [('h', '5 Первый этап и демонстрация'),
                   ('p',
                    'Окно разработки — 4–9 октября 2026 года. По официальному уведомлению университетский отбор завершается к 10 октября; внутренний срок '
                    'подачи нужно сверить с вузом. Регистрация и оплата регионального этапа заканчиваются 15 октября в 20:00. Доработка к региональному и '
                    'национальному этапам проводится отдельно, без переноса текущей сдачи на декабрь.'),
                   ('table',
                    [['Дата', 'Участник A', 'Участник B', 'Результат'],
                     ['4–5 октября', 'Контракт и текстовый API с сохранением', 'Источники и примеры; каркас интерфейса', 'Запускаемый каркас и версии данных'],
                     ['6 октября', 'Модель, проверки, цитаты и API', 'Аудио и ASR; интерфейс', 'Текстовый черновик и аудиоввод'],
                     ['7 октября', 'Сохранение, подтверждение, экспорт; интеграция', 'Экран сверки цитат и аудио; интеграция', 'Три сквозных сценария'],
                     ['8 октября',
                      'Измерения и реальные результаты',
                      'Повторная проверка, восстановление после ошибки и работа без сети',
                      'Таблица и список ошибок'],
                     ['9 октября', 'Китайские материалы и пакет подачи', 'Китайское видео и проверка запуска из пакета', 'Видео и материалы']]),
                   ('h2', 'Показ проекта'),
                   ('p',
                    'Ориентир видео — 4 минуты. Обычный случай показывает китайский ввод, поля и доказательства; аллергия — проверку правил; неполный случай — '
                    'неизвестность и исправление врачом. В конце показаны подтверждение, экспорт и фактические результаты. Монтаж не должен скрывать сбои '
                    'модели или задержку.'),
                   ('h2', 'Материалы и ресурсы'),
                   ('p',
                    'Пакет включает китайский PDF технического предложения, китайское MP4, PDF презентации, код, инструкции, версии моделей, перечень данных и '
                    'скрипты оценки. Подтверждения сотрудничества, патентов и пилотов прилагаются только при их наличии.'),
                   ('p',
                    'Сначала тестируется последовательная обработка на устройстве с 16 ГБ памяти. Замеряются пиковая память после загрузки и задержка; при '
                    'необходимости уменьшаются модель и длина входа. Покупка GPU или аренда обучения не заложены как обязательное условие. Инструкция '
                    'фиксирует версии, загрузку моделей и обработку ошибок.')],
                  [('h', '6 Польза риски и развитие'),
                   ('p', 'Ожидаемая польза — помощь в организации содержания разговора и проверке оснований записи.'),
                   ('p',
                    'Экономия времени, уменьшение пропусков и пригодность для клиники требуют дальнейших пользовательских и медицинских проверок. Моделируемые '
                    'тесты не заменяют клиническое исследование и не доказывают безопасность диагностики или лекарств.'),
                   ('p',
                    'Начальный путь применения — учебные и моделируемые записи. После проверки китайскими медицинскими специалистами и контролируемого '
                    'испытания можно обсуждать больничные интерфейсы. Возможный сервис развёртывания и поддержки остаётся гипотезой; тарифы, закупки и выручка '
                    'не подтверждены.'),
                   ('h2', 'Риски и условия расширения'),
                   ('p',
                    'Китайская медицинская лексика требует оригинальных цитат, перекрёстной проверки и китайского рецензирования. ASR может изменить отрицание '
                    'или дозу, поэтому важные поля возвращают пользователя к исходнику. Ограниченные правила явно показывают неизвестность. Источники с '
                    'неясной лицензией не включаются в официальный набор; реальные данные потребуют отдельного разрешения и процесса защиты.'),
                   ('p',
                    'После этапа порядок улучшений определяется частотой ошибок: сначала данные и подсказки, затем расширение поиска. Дообучение '
                    'рассматривается при наличии достаточно большого разрешённого корпуса с разметкой полей, завершённого baseline и подходящих ресурсов. '
                    'Масштабное тестирование, надёжные роли говорящих и реальные интеграции — последующие задачи.'),
                   ('h2', 'Источники'),
                   ('ref',
                    'Официальное уведомление и требования к материалам\n'
                    'https://www.aicomp.cn/notice/notice-1/3674.html\n'
                    'https://www.aicomp.cn/tracks/tracks-2/3775.html'),
                   ('ref', 'IMCS-21 Данные и исследовательская реализация\nhttps://github.com/lemuria-wchen/imcs21'),
                   ('ref',
                    'Дополнительные кандидаты DISC-MedLLM и Huatuo-26M\n'
                    'https://github.com/FudanDISC/DISC-MedLLM\n'
                    'https://github.com/FreedomIntelligence/Huatuo-26M'),
                   ('ref', 'Текущая модель и планируемый речевой инструмент\nhttps://huggingface.co/Qwen/Qwen3.5-4B\nhttps://github.com/modelscope/FunASR'),
                   ('p', 'Источники проверены 4 октября 2026 года. Включение данных зависит от фактической проверки условий и образцов.')]]},
 'EN': {'title': 'Outpatient clinical documentation assistant',
        'subtitle': 'SmartMed Technical proposal',
        'meta': 'Eighth Global Campus Artificial Intelligence Algorithm Elite Competition\nAlgorithm Innovation Track Topic 4 AI+场景创新\n5 October 2026',
        'pages': [[('h', '1 Project overview'),
                   ('h2', 'Application summary'),
                   ('p',
                    'SmartMed is proposed as an assistant that converts Chinese clinician–patient dialogue into a traceable clinical note draft. Chinese '
                    'speech recognition, an existing language model, terminology retrieval and deterministic checks will connect key fields to source '
                    'statements, leave missing facts unknown, and flag allergies, duplicate ingredients and incomplete medication information. Export follows '
                    'clinician review. The first stage targets a local Chinese demonstration and independent tests of completeness, factual consistency and '
                    'workflow reliability. Clinical effectiveness remains to be validated.'),
                   ('h2', 'Scenario and user need'),
                   ('p',
                    'The intended user is a clinician preparing a note after an outpatient consultation. Symptoms, duration, negatives, history and '
                    'medications are distributed across turns. The proposed screen places dialogue and audio segments alongside structured fields, source '
                    'evidence and items requiring clarification.'),
                   ('p',
                    'The first stage uses simulated Chinese consultations. Text input is validated before short audio files are added. SOAP organizes '
                    'subjective information, objective observations, assessment already expressed by the clinician, and the stated plan. Diagnoses, '
                    'examination results and treatment absent from the input must not be invented.'),
                   ('h2', 'First stage objective'),
                   ('p',
                    'As of 5 October, local Qwen3.5 4B extraction, Instructor and Pydantic validation, SQLite revisions, editing, confirmation and JSON export '
                    'are implemented. Four Qwen models have been compared and known failures regression-tested. The API extracts six facts; full SOAP, the '
                    'audio interface, terminology retrieval and medication checks remain development work. There is no clinical pilot or medical validation.'),
                   ('h2', 'Team capability and validation boundary'),
                   ('p',
                    'The team comprises two software developers. Member A handles the backend, model, field contract, validation and storage. Member B handles '
                    'audio, interface, examples and presentation materials; cases are cross-checked. No clinician or clinical adviser is on the team. '
                    'Developers operate simulated cases; confirmation is not clinician validation.')],
                  [('h', '2 Architecture and processing'),
                   ('p',
                    'The proposed system has one Chinese web interface and one backend. React and TypeScript serve the interface, FastAPI provides the '
                    'service, and SQLite stores jobs, draft versions and confirmation records. Audio and data files remain local. The backend uses explicit '
                    'workflow states and model adapters.'),
                   ('h2', 'Input to reviewable draft'),
                   ('p',
                    'Chinese text or audio → transcription and correction → structured extraction → evidence links → terminology candidates → rules → '
                    'clinician editing and confirmation → export. Text bypasses speech recognition. Audio first produces a transcript for correction. Failures '
                    'retain the job and display a specific reason.'),
                   ('p',
                    'The speech component will initially test a Chinese FunASR model. Automatic speaker roles are optional; unreliable separation requires '
                    'manual labels. Models and dictionaries are downloaded before an offline demonstration.'),
                   ('p',
                    'The current extractor is locally quantized Qwen3.5 4B with thinking output disabled. Controls are Qwen2.5 3B, Qwen3 4B Instruct and '
                    'Qwen2.5 7B. Instructor passes a JSON Schema restricting statuses, values and evidence. Citations are selected from original patient spans '
                    'and checked against the transcript. One retry is allowed before failure. Structural validity does not establish medical correctness.'),
                   ('h2', 'Field contract and clinician control'),
                   ('p',
                    'Current fields are fever, cough, duration, penicillin allergy, medication taken and dose taken. Status is present, negated or unknown; '
                    'unresolved conflict is unknown. Values are copied without normalization of numbers or units, and negation also needs evidence. The '
                    'clinician is the intended user role; developers operate that role in simulated demonstrations.'),
                   ('p',
                    'Diagnosis terminology retrieval is triggered by content already expressed by the clinician. Candidates include dictionary source and '
                    'version. Without a verifiable coding source, the system shows terms without fabricated ICD codes. Initial rules cover known allergy '
                    'matching, duplicate ingredients and incomplete dosage information. Outside coverage or with insufficient information, the result is '
                    'unknown rather than safe.'),
                   ('p',
                    'The clinician edits fields, inspects evidence and resolves warnings. Confirmation enables JSON and one printable format, recording the '
                    'version and time. This is a prototype action record without a claim of legally valid electronic signature.')],
                  [('h', '3 Data and provenance'),
                   ('h2', 'Sources before generation'),
                   ('p',
                    'Four data types are separated: dialogues, reference field annotations, speech tests and rule dictionaries. Provenance, permissions, '
                    'availability and quality are checked before inclusion. Medical question answering is not treated as ready-made SOAP supervision. Each '
                    'case records its source, original ID, version, transformation and split. Variants of one source cannot cross development and test sets.'),
                   ('p',
                    'IMCS-21 is the first candidate: Chinese consultations with a medical report generation task. Its primarily pediatric scope and report '
                    'format differ from the proposed SOAP contract. Data conditions must be checked and field-to-turn mappings annotated manually. Published '
                    'benchmark results are not SmartMed results.'),
                   ('p',
                    'DISC-Med-SFT is a supplementary candidate with processed and reconstructed consultations. Huatuo-26M mainly contains medical questions '
                    'and answers and may support terminology and scenario exploration. Neither provides real consultation audio or an independent project gold '
                    'standard. A code license does not establish permission for every underlying source.'),
                   ('h2', 'Small first stage evaluation set'),
                   ('p',
                    'Four-model comparison used 40 AI-assisted synthetic Chinese dialogues and 240 fields, with no real patient data. References were not '
                    'reviewed by a clinician or native Chinese speaker. This set served model selection, and some cases subsequently informed prompt and '
                    'validator development; it is not an independent final test. A new frozen set outside development is required, with provenance and review '
                    'records.'),
                   ('p',
                    'Annotations cover symptoms, time, negatives, allergies, medications and explicitly stated doses, with source turns. Both members '
                    'cross-check annotations and record disagreements. A fluent Chinese reviewer is sought when available; clinical interpretation requires '
                    'medical expertise. Machine translation is not medical review.'),
                   ('h2', 'Audio and dictionaries'),
                   ('p',
                    'Six simulated Chinese clips of 1–2 minutes are planned, split equally between development and testing, with manual transcripts and roles. '
                    'Synthetic speech and human readings are labeled separately; they do not represent real clinical acoustics. General speech corpora can '
                    'provide supplementary tests only.'),
                   ('p',
                    'Medication ingredients, allergy mappings and diagnosis terms use small, traceable versioned dictionaries. Each rule records its source, '
                    'conditions and coverage. Dose ceilings and complex interactions are excluded without reliable support. No real patient recordings are '
                    'collected in the first stage.')],
                  [('h', '4 Contribution and evaluation'),
                   ('h2', 'Testable scenario contributions'),
                   ('p',
                    'First, each note field links to dialogue evidence on the same screen. Second, negatives, unknowns and conflicts are explicit states that '
                    'discourage filling gaps by inference. Third, terminology retrieval and deterministic checks follow extraction, with coverage and '
                    'clinician confirmation integrated into the workflow.'),
                   ('p',
                    'These are proposed engineering contributions requiring validation. Original foundation model research and superiority over clinical '
                    'products are not claimed. External models, tools and data are credited. The team contribution concerns workflow, contracts, interface, '
                    'rules and evaluation.'),
                   ('h2', 'Comparison and component analysis'),
                   ('p',
                    'The original comparison kept 40 cases, six fields, prompt and configuration fixed. Qwen3.5 4B achieved 235/240 correct statuses and 28/40 '
                    'structurally valid cases; Qwen3 4B Instruct achieved 230/240 and 37/40. Qwen3.5 median latency over 14 warm requests was 7.73 seconds, '
                    'under a different configuration from the current API. After revisions, five known cases matched reference statuses and values, and 53 '
                    'backend tests passed. Regression does not establish clinical reliability or independent-test improvement. Further experiments require a '
                    'new frozen set.'),
                   ('p',
                    'Audio evaluation compares extraction from manual transcripts with extraction from ASR transcripts to separate recognition and extraction '
                    'errors. Per-case errors, latency and resource use are retained. Findings from this small simulated set apply only to the tested '
                    'scenarios.'),
                   ('table',
                    [['Evaluation item', 'Method'],
                     ['Field extraction', 'P, R and F1 using predefined field/value matching; negatives and doses reported separately.'],
                     ['Facts and evidence', 'Manual counts of unsupported facts, correct evidence links and correctly retained unknowns.'],
                     ['Dictionary and rules', 'Candidate hits and warning correctness within coverage; uncovered and insufficient cases reported separately.'],
                     ['Speech and workflow', 'CER against manual transcripts; DER only with reference roles; confirmation and export checks.']]),
                   ('h2', 'Acceptance conditions'),
                   ('p',
                    'Three Chinese scenarios must complete the workflow: normal documentation, allergy and missing information. Invalid JSON cannot reach '
                    'confirmation; absent doses remain unknown in targeted tests; unconfirmed drafts cannot be exported as confirmed notes. Numerical results '
                    'are filled after experiments. No F1, speed or time-saving percentage is promised in advance.')],
                  [('h', '5 First stage implementation and demonstration'),
                   ('p',
                    'The development window is 4–9 October 2026. The official notice sets campus selection completion by 10 October; the internal submission '
                    'time must be checked with the institution. Regional registration and payment close at 20:00 on 15 October. Regional and national '
                    'preparations are subsequent iterations, rather than a December deadline for the current submission.'),
                   ('table',
                    [['Date', 'Member A', 'Member B', 'Deliverable'],
                     ['4–5 October',
                      'Field contract, text API and storage',
                      'Data sources, cases and interface scaffold',
                      'Runnable scaffold and versioned data'],
                     ['6 October', 'Model, validation, evidence and API', 'Audio, ASR and interface', 'Text draft and audio input'],
                     ['7 October',
                      'Storage, confirmation, export and integration',
                      'Evidence and audio review interface; integration',
                      'Three complete scenarios'],
                     ['8 October', 'Experiments and actual results', 'Repeat tests, failure recovery and offline run', 'Results and error list'],
                     ['9 October',
                      'Chinese proposal, presentation content and submission package',
                      'Chinese video and packaged launch verification',
                      'Video and materials']]),
                   ('h2', 'Demonstration'),
                   ('p',
                    'The target video length is about four minutes. A normal case shows Chinese input, fields and evidence; an allergy case shows rules; an '
                    'incomplete case shows unknown values and clinician correction. Confirmation, export and actual evaluation results conclude the '
                    'demonstration. Editing must not conceal model failures or latency.'),
                   ('h2', 'Materials and resources'),
                   ('p',
                    'The package includes a Chinese technical proposal PDF, Chinese MP4, presentation exported to PDF, code, instructions, model versions, a '
                    'data manifest and evaluation scripts. Partnership, patent and pilot evidence is supplied only when genuine and available.'),
                   ('p',
                    'Sequential processing is initially tested on a device with 16 GB memory. Peak memory after model loading and latency are measured; model '
                    'size and input length can be reduced. Purchasing GPUs or renting training resources is not a prerequisite. Launch instructions record '
                    'versions, model downloads and failure handling.')],
                  [('h', '6 Value risks and development'),
                   ('p',
                    'Expected value is assistance with organizing consultation content and checking the basis of a note. Time savings, fewer omissions and '
                    'clinical suitability require subsequent user studies and medical validation. Simulated tests do not replace clinical research or '
                    'establish diagnostic and medication safety.'),
                   ('p',
                    'An initial deployment path is educational and simulated documentation, followed by Chinese medical review and controlled trials before '
                    'hospital interfaces are considered. Deployment and maintenance services are a possible commercial hypothesis; pricing, procurement and '
                    'revenue are unvalidated.'),
                   ('h2', 'Risks and extension conditions'),
                   ('p',
                    'Chinese medical expressions require source quotations, cross-checking and Chinese review. ASR may alter a negative or dose, so critical '
                    'fields must link back to the input. Limited rules explicitly show unknown states. Sources with unclear permissions do not enter the '
                    'formal dataset; real data will need a separate authorization and protection process.'),
                   ('p',
                    'After the first stage, error frequency determines priorities: improve data and prompts first, then expand retrieval. Fine tuning is '
                    'considered only with sufficient permitted dialogue and field annotations, a completed baseline and suitable resources. Larger tests, '
                    'reliable speaker roles and real integrations remain future work.'),
                   ('h2', 'References'),
                   ('ref',
                    'Official notice and material requirements\n'
                    'https://www.aicomp.cn/notice/notice-1/3674.html\n'
                    'https://www.aicomp.cn/tracks/tracks-2/3775.html'),
                   ('ref', 'IMCS-21 Data and research implementation\nhttps://github.com/lemuria-wchen/imcs21'),
                   ('ref',
                    'Supplementary candidates DISC-MedLLM and Huatuo-26M\n'
                    'https://github.com/FudanDISC/DISC-MedLLM\n'
                    'https://github.com/FreedomIntelligence/Huatuo-26M'),
                   ('ref', 'Current model and planned speech tool\nhttps://huggingface.co/Qwen/Qwen3.5-4B\nhttps://github.com/modelscope/FunASR'),
                   ('p', 'Sources checked on 4 October 2026. Dataset inclusion remains subject to permission and sample review.')]]}}

def font(style, size, lang):
    style.font.name = 'Noto Sans CJK SC' if lang == 'ZH' else 'Arial'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.italic = False
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    for key in list(fonts.attrib):
        if 'Theme' in key:
            del fonts.attrib[key]
    color = style.element.get_or_add_rPr().find(qn('w:color'))
    if color is not None:
        for key in list(color.attrib):
            if 'theme' in key:
                del color.attrib[key]

def make(lang, content):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin, s.bottom_margin = Cm(1.9), Cm(1.8)
    s.left_margin, s.right_margin = Cm(2.0), Cm(2.0)
    for name, size in [('Normal', 10.5), ('Title', 22), ('Subtitle',12), ('Heading 1',16), ('Heading 2',12)]:
        font(d.styles[name], size, lang)
    for style in d.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    normal = d.styles['Normal'].paragraph_format
    normal.space_after = Pt(7)
    normal.line_spacing = 1.12
    for name in ['Heading 1','Heading 2']:
        p = d.styles[name].paragraph_format
        p.space_before, p.space_after = Pt(11), Pt(6)
        p.keep_with_next = True
    d.core_properties.author = ''
    d.core_properties.last_modified_by = ''
    d.core_properties.title = content['title']
    d.core_properties.subject = 'SmartMed technical proposal'
    d.add_paragraph(content['title'], 'Title')
    d.add_paragraph(content['subtitle'], 'Subtitle')
    d.add_paragraph(content['meta'])
    for idx, page in enumerate(content['pages']):
        for position, (kind, text) in enumerate(page):
            if kind == 'h':
                paragraph = d.add_paragraph(text, 'Heading 1')
                if idx and position == 0:
                    paragraph.paragraph_format.page_break_before = True
            elif kind == 'h2':
                d.add_paragraph(text, 'Heading 2')
            elif kind == 'table':
                t = d.add_table(rows=0, cols=len(text[0]))
                t.autofit = False
                total = 17.0
                widths = [total/len(text[0])] * len(text[0])
                if len(text[0]) == 4:
                    widths = [2.1,5.05,5.05,4.8]
                else:
                    widths = [3.5,13.5]
                for i, width in enumerate(widths):
                    t.columns[i].width = Cm(width)
                for rowid, values in enumerate(text):
                    row = t.add_row()
                    trpr = row._tr.get_or_add_trPr()
                    no_split = OxmlElement('w:cantSplit')
                    trpr.append(no_split)
                    if rowid == 0:
                        trpr.append(OxmlElement('w:tblHeader'))
                    for i, value in enumerate(values):
                        c = row.cells[i]
                        c.width = Cm(widths[i])
                        c.text = value
                        props = c._tc.get_or_add_tcPr()
                        margins = OxmlElement('w:tcMar')
                        for side in ('top','left','bottom','right'):
                            el = OxmlElement('w:'+side)
                            el.set(qn('w:w'), '85')
                            el.set(qn('w:type'),'dxa')
                            margins.append(el)
                        props.append(margins)
                        borders = OxmlElement('w:tcBorders')
                        for side in ('top','left','bottom','right'):
                            el = OxmlElement('w:'+side)
                            el.set(qn('w:val'),'single')
                            el.set(qn('w:sz'),'4')
                            el.set(qn('w:color'),'D2D2D2')
                            borders.append(el)
                        props.append(borders)
                        if rowid == 0:
                            shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'F2F2F2'); props.append(shd)
                        for p in c.paragraphs:
                            p.paragraph_format.space_after=Pt(2)
                            p.paragraph_format.line_spacing=1.06
                            for r in p.runs:
                                r.font.size=Pt(9.5)
                                r.bold=rowid == 0
                d.add_paragraph().paragraph_format.space_after=Pt(0)
            else:
                p = d.add_paragraph(text)
                if kind == 'ref':
                    p.paragraph_format.keep_together=True
                    for r in p.runs:
                        r.font.size=Pt(9)
    footer=s.footer.paragraphs[0]
    footer.alignment=2
    r=footer.add_run()
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); r._r.addnext(fld)
    output=ROOT/f'SmartMed_Ambient_EHR_Copilot_{lang}.docx'
    d.save(output)
    print(output)

if __name__ == '__main__':
    for lang, content in CONTENT.items():
        make(lang,content)
